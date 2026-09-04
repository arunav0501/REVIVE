#!/usr/bin/env python3
"""
Generate a professional, fully-styled Word Document (.docx)
containing the complete architectural, mathematical, and operational guide for RecoverAI.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def style_table(table, header_bg="1E3A8A", alt_bg="F8FAFC", border_color="CBD5E1"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Set thin subtle borders
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

    # Style header row
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

    # Style alternating body rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(31, 41, 55)

def add_callout(doc, title, text, bg_color="EFF6FF", border_color="3B82F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:bottom w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.font.bold = True
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(30, 58, 138)

    run_b = p.add_run(text)
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_document():
    doc = Document()

    # Configure Margins (1 inch all around)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Base Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(31, 41, 55)

    # -------------------------------------------------------------
    # DOCUMENT COVER / TITLE HEADER
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("RecoverAI — Autonomous AR Follow-Up & Revenue Cycle AI Platform")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)  # Navy Blue

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Comprehensive System Architecture, Machine Learning Models, Vector RAG Knowledge Base, and Operational Guide")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata Strip
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(20)
    r1 = p_meta.add_run("Document Version: ")
    r1.font.bold = True
    p_meta.add_run("1.0.0  |  ")
    r2 = p_meta.add_run("Platform Engine: ")
    r2.font.bold = True
    p_meta.add_run("FastAPI + React 19 + Scikit-Learn  |  ")
    r3 = p_meta.add_run("Target Audience: ")
    r3.font.bold = True
    p_meta.add_run("Executives, RCM Directors, Engineers & Data Scientists")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # SECTION 1: EXECUTIVE SUMMARY & PROBLEM CONTEXT
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. Executive Summary & Healthcare Revenue Context", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "RecoverAI is an autonomous artificial intelligence platform engineered specifically for healthcare revenue cycle management (RCM), "
        "hospital billing operations, and physician practices. The platform eliminates the manual bottlenecks of insurance claim follow-ups, "
        "denial dispute packet creation, and revenue recovery."
    )

    add_callout(
        doc,
        "The $260 Billion Problem: Healthcare Claim Denials & AR Aging",
        "• $260+ Billion in Annual Denials: Approximately 15% to 20% of all medical claims submitted to commercial and government insurance payers are initially denied or delayed.\n"
        "• Revenue Leakage & Aging: Claims languish in Accounts Receivable (AR) past 30, 60, 90, or 120+ days. Without timely clinical appeals, these claims exceed strict timely filing limits and are written off as uncollectible bad debt.\n"
        "• Labor-Intensive Workflow: Hospital staff spend hours manually logging into payer portals, waiting on phone hold queues, and writing manual appeals.\n"
        "• Arbitrary Prioritization: Human teams work claims randomly rather than optimizing for mathematically expected dollar recovery."
    )

    doc.add_paragraph(
        "RecoverAI solves these challenges by combining dual machine learning classifiers, a financially-aware priority yield algorithm, "
        "a grounded TF-IDF vector policy retrieval engine (RAG), a multi-format communication synthesizer, a 4-class adjudication simulator, "
        "and an autonomous closed-loop execution pipeline."
    )

    # -------------------------------------------------------------
    # SECTION 2: END-TO-END SYSTEM ARCHITECTURE
    # -------------------------------------------------------------
    h1 = doc.add_heading("2. High-Level Architecture & End-to-End Data Pipeline", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The platform is structured into five cohesive layers operating in real time:"
    )

    arch_table = doc.add_table(rows=6, cols=3)
    arch_table.rows[0].cells[0].text = "Layer"
    arch_table.rows[0].cells[1].text = "Components / Modules"
    arch_table.rows[0].cells[2].text = "Primary Operational Role"

    arch_data = [
        ("Data & Storage Layer", "SQLite (ar_followup.db), SQLAlchemy 2.0 ORM, 10,000 Claims Corpus", "Maintains 10,000 realistic claims, 10 payer policy profiles, 10,109 follow-up records, and longitudinal recovery outcomes."),
        ("Intelligence & ML Layer", "Dual HistGradientBoosting Models, Priority Scoring Engine", "Predicts P(Delay) and P(Denial), calculates transparent P(Recovery), and assigns 0-100 Priority Scores into Critical/High/Med/Low bands."),
        ("Vector RAG Policy Layer", "TfidfVectorizer, Cosine Similarity Index, 40 Policy Chunks", "Indexes payer timely filing windows, mandatory clinical documentation requirements, dispute steps, and contact protocols."),
        ("Autonomous Agent & Generator", "AIFollowUpAgent, Multi-Format Artifact Generator", "Formulates optimal action (Appeal, Inquiry, Resubmission, Escalation) and drafts formal CMS-1500/UB-04 letters, portal notes, and phone scripts."),
        ("Simulation & Closed-Loop", "PayerResponseSimulator, ClosedLoopRecoveryEngine", "Executes 4-class adjudication simulation (Approved Full, Partial, Additional Info, Denied), applies database ledger updates, and logs recoveries.")
    ]

    for idx, (l, c, r) in enumerate(arch_data, start=1):
        arch_table.rows[idx].cells[0].text = l
        arch_table.rows[idx].cells[1].text = c
        arch_table.rows[idx].cells[2].text = r

    style_table(arch_table)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 3: MACHINE LEARNING MODELS & MATHEMATICAL FORMULAS
    # -------------------------------------------------------------
    h1 = doc.add_heading("3. Machine Learning Models & Mathematical Formulations", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "RecoverAI deploys dual Scikit-Learn gradient boosting models alongside a transparent multi-variable prioritization engine."
    )

    # ML Benchmarks Table
    ml_table = doc.add_table(rows=8, cols=3)
    ml_table.rows[0].cells[0].text = "Model Metric / Attribute"
    ml_table.rows[0].cells[1].text = "Adjudication Delay Classifier"
    ml_table.rows[0].cells[2].text = "Claim Denial Risk Classifier"

    ml_data = [
        ("Target Variable", "was_delayed (Binary: 0 or 1)", "was_denied (Binary: 0 or 1)"),
        ("Algorithm", "HistGradientBoostingClassifier", "HistGradientBoostingClassifier"),
        ("ROC-AUC Score", "0.7462", "0.7818"),
        ("Accuracy", "79.20%", "82.00%"),
        ("Precision", "82.07%", "68.06%"),
        ("Recall (Sensitivity)", "93.05%", "54.64%"),
        ("F1-Score", "0.8722", "0.6061"),
    ]

    for idx, (m, d, dn) in enumerate(ml_data, start=1):
        ml_table.rows[idx].cells[0].text = m
        ml_table.rows[idx].cells[1].text = d
        ml_table.rows[idx].cells[2].text = dn

    style_table(ml_table)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Feature Vector Breakdown
    h2 = doc.add_heading("3.1 24-Dimensional Input Feature Vector", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "The models evaluate 24 distinct clinical, financial, and administrative features:\n"
        "• 13 Numeric Features: days_in_ar, outstanding_amount, billed_amount, allowed_amount, paid_amount, "
        "patient_responsibility, expected_reimbursement, followup_count, resubmission_count, days_since_last_response, "
        "avg_processing_days, avg_response_days, filing_deadline_days.\n"
        "• 11 Categorical Features: payer_id, claim_type, place_of_service, claim_status, denial_status, "
        "denial_code, denial_reason, authorization_required, authorization_status, network_status, appeal_status."
    )

    # Formulas
    h2 = doc.add_heading("3.2 Core Prioritization & Financial Yield Formulas", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    add_callout(
        doc,
        "Mathematical Formulas for Prioritization & Recovery",
        "1. Base Recovery Probability:\n"
        "   Base_Recovery = max(0.05, 1.0 - 0.65 * P(Denial) - 0.15 * P(Delay))\n\n"
        "2. Clinical & Aging Adjustments:\n"
        "   • Out-of-Network: Multiply by 0.85\n"
        "   • Prior Authorization Denied: Multiply by 0.60 (Pending: 0.88, Approved: min(0.98, Base * 1.08))\n"
        "   • Appeal Under Review: min(0.95, Base * 1.05)\n"
        "   • AR Aging Decay: If Days > 120 (0.82x); If Days > 90 (0.90x)\n"
        "   • Final P(Recovery) is clamped to [0.05, 0.98]\n\n"
        "3. Expected Financial Recovery ($):\n"
        "   Expected_Recovery = Outstanding_Amount * P(Recovery)\n\n"
        "4. Priority Score (0 - 100):\n"
        "   Priority_Score = [ 0.45 * P(Recovery) + 0.30 * P(Delay) + 0.25 * min(1.0, Outstanding / $50,000) ] * 100\n\n"
        "5. Priority Bands:\n"
        "   • Critical: 85.0 - 100.0 (Immediate expedited appeal)\n"
        "   • High: 70.0 - 84.9 (24-hour follow-up SLA)\n"
        "   • Medium: 40.0 - 69.9 (Routine automated inquiry)\n"
        "   • Low: 0.0 - 39.9 (Batch automated monitoring)"
    )

    # -------------------------------------------------------------
    # SECTION 4: VECTOR RAG ENGINE
    # -------------------------------------------------------------
    h1 = doc.add_heading("4. Retrieval-Augmented Generation (RAG) Architecture", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "In healthcare billing, generative hallucinations lead to legal dismissal of appeals. "
        "RecoverAI's vector RAG engine grounds every decision and appeal letter in verified contractual policies."
    )

    doc.add_paragraph(
        "• Corpus: 40 structured policy documents across 10 major commercial and government payers "
        "(Aetna, Blue Cross Blue Shield, UnitedHealthcare, Cigna, Humana, Medicare Part A, Medicare Part B, Medicaid State, Kaiser Permanente, Anthem BCBS).\n"
        "• 4 Specialized Chunks Per Payer:\n"
        "   1. Timely Filing, Adjudication & Escalation Limits (e.g. 90-day submission limit, 15-day processing window).\n"
        "   2. Prior Authorization & Mandatory Clinical Documentation (e.g. operative notes, pre-cert certificates).\n"
        "   3. Appeals, Disputes & Peer-to-Peer Guidelines (e.g. medical necessity disputes, CO-16/CO-50 dispute protocols).\n"
        "   4. Communication Channels & Contact Protocols (e.g. Availity portal, EDI 276/277 transactions, certified mail).\n"
        "• Vector Index: Scikit-Learn TfidfVectorizer with ngram_range=(1,2) and sublinear term-frequency scaling.\n"
        "• Payer Metadata Boosting: Inquiries automatically boost matching payer chunks with a +0.15 similarity weight.\n"
        "• Dynamic Claim Context Retrieval: When drafting an appeal, a targeted query like 'Inpatient denial CO-16 missing documentation appeal' "
        "is executed to extract the exact contract paragraphs for citation."
    )

    # -------------------------------------------------------------
    # SECTION 5: COMPLETE TAB-BY-TAB WALKTHROUGH
    # -------------------------------------------------------------
    h1 = doc.add_heading("5. Comprehensive Walkthrough of All 10 Frontend Tabs", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    # Tab 1
    doc.add_heading("5.1 Tab 1: Executive Dashboard (Mission Control)", level=2)
    doc.add_paragraph(
        "The executive command center provides real-time financial oversight across the entire hospital AR portfolio:\n"
        "• Total AR Balance: Real-time sum of all outstanding balances across 10,000 claims (~$44.05M).\n"
        "• Total Claims: 10,000 active hospital records.\n"
        "• Critical Priority Claims: Claims scoring >= 85 requiring immediate intervention (~1,200 - 1,400 claims).\n"
        "• Expected Financial Recovery: Total yield calculated via ML (~$28M - $30M).\n"
        "• Recovered Revenue: Cumulative cash recovered and ledgered via the autonomous loop.\n"
        "• Average Days in AR: Mean age of unpaid claims (~52.4 days).\n"
        "• Denial Rate: Proportion of portfolio claims currently denied (~25.4%).\n"
        "• Visual Charts: AR Aging Intervals (0-30, 31-60, 61-90, 91-120, 120+ days) vs Expected Recovery; "
        "Priority Band Distribution Donut; Top Denial Root Causes (CO-16, CO-50, CO-197); Payer Exposure Table."
    )

    # Tab 2
    doc.add_heading("5.2 Tab 2: All Claims (Claim Portfolio Explorer)", level=2)
    doc.add_paragraph(
        "Interactive claim workspace allowing granular inspection of any claim:\n"
        "• Multi-Parameter Filters: Search by Claim ID, Patient ID, Provider ID, Denial text; filter by Payer, Status, and Priority Band.\n"
        "• 4-Tab Slide-Out Detail Modal:\n"
        "   1. AI Strategy Workspace: Visual radial gauges for Delay Risk, Denial Probability, Recovery Potential, and Priority Score; deterministic explainability factor bullet points.\n"
        "   2. Clinical & Financial Ledger: Line-item breakdown of Billed, Allowed, Paid, Patient Responsibility, and Outstanding balances, plus POS, CPT, and ICD-10 codes.\n"
        "   3. Payer Rules & RAG Citations: Active policy paragraphs citing timely filing limits and required documentation.\n"
        "   4. Longitudinal Timeline: Historical audit trail of past follow-up attempts, EDI 276 inquiries, and remittance advice."
    )

    # Tab 3
    doc.add_heading("5.3 Tab 3: Priority Worklist (High-Yield Queue)", level=2)
    doc.add_paragraph(
        "A noise-free operational queue displaying exclusively Critical (85+) and High (70-84.9) priority claims, sorted strictly by expected financial recovery yield."
    )

    # Tab 4
    doc.add_heading("5.4 Tab 4: Closed-Loop Engine (Autonomous Batch Recovery)", level=2)
    doc.add_paragraph(
        "The autonomous execution pipeline that processes claims in unattended batches:\n"
        "• Controls: Batch Size (10, 20, 30, 50), Minimum Priority Band, Payer Filter, Action Quality.\n"
        "• 7-Step Cycle: Ingestion -> AI Strategy -> RAG Grounding -> Artifact Synthesis -> Adjudication Simulation -> Database Ledger Update -> Analytics Recomputation.\n"
        "• Output: Live progress bar, real-time logs, recovery summary banner (Disputed $, Recovered $, Recovery Rate %), and line-item transaction ledger with EFT reference codes."
    )

    # Tab 5
    doc.add_heading("5.5 Tab 5: Payer Response Simulator (Adjudication Sandbox)", level=2)
    doc.add_paragraph(
        "Probabilistic simulation sandbox to model how payers adjudicate follow-up actions:\n"
        "• Configurable Parameters: Claim selection, Action Type (Appeal, Status Inquiry, Resubmission, Escalation), Channel, Quality.\n"
        "• 4-Class Probability Output: P(APPROVED_FULL), P(APPROVED_PARTIAL), P(ADDITIONAL_INFO_REQUIRED), P(DENIAL_UPHELD).\n"
        "• Remittance Generation: Generates ANSI 835 Remittance Advice text, EFT reference numbers (e.g. EFT-UHC-849201), and updates SQLite balances."
    )

    # Tab 6
    doc.add_heading("5.6 Tab 6: Multi-Format Generator Studio", level=2)
    doc.add_paragraph(
        "Generates tailored legal and clinical communication artifacts:\n"
        "• 5 Supported Formats:\n"
        "   1. Formal CMS-1500 / UB-04 Appeal Letters (complete with grounded policy citations and exhibit checklists).\n"
        "   2. Payer Portal Messages (optimized for Availity, NaviNet, Optum, Change Healthcare).\n"
        "   3. Live Call Scripts (with IVR prompts, verification checkpoints, and Supervisor Escalation Triggers).\n"
        "   4. Secure Provider Email Templates (HIPAA-compliant de-identified summaries).\n"
        "   5. Electronic Resubmission Notes (ANSI 837 replacement notes with modifier -59/-76 explanations).\n"
        "• Tone Selector: Firm, Urgent, Collaborative; Custom Notes input; One-Click Clipboard Copy."
    )

    # Tab 7
    doc.add_heading("5.7 Tab 7: Autonomous AI Agent (Decision Formulation)", level=2)
    doc.add_paragraph(
        "Single-claim AI reasoning workspace displaying Chain-of-Thought reasoning, confidence scores (85%-95%), urgency rating, and one-click 'Execute & Record in Ledger' button."
    )

    # Tab 8
    doc.add_heading("5.8 Tab 8: Payer Policy Knowledge Base (RAG Search)", level=2)
    doc.add_paragraph(
        "Live semantic search explorer over the 40 policy chunks with similarity score breakdown and metadata filters."
    )

    # Tab 9
    doc.add_heading("5.9 Tab 9: ML Model Diagnostics & Benchmarks", level=2)
    doc.add_paragraph(
        "Technical metrics dashboard displaying ROC-AUC curves, accuracy, precision, recall, F1-scores, and 24-feature schema specifications."
    )

    # Tab 10
    doc.add_heading("5.10 Tab 10: System Diagnostics & Health Status", level=2)
    doc.add_paragraph(
        "System telemetry monitoring FastAPI health status, SQLite database URI, and live counts for 10,000 Claims, 10 Payers, 10 Policies, 10,109 Follow-ups, and 10,000 Outcomes."
    )

    # -------------------------------------------------------------
    # SECTION 6: DENIAL CODES & DECISION MATRIX
    # -------------------------------------------------------------
    h1 = doc.add_heading("6. CARC Denial Codes & Action Matrix Reference", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    carc_table = doc.add_table(rows=7, cols=4)
    carc_table.rows[0].cells[0].text = "CARC Code"
    carc_table.rows[0].cells[1].text = "Description"
    carc_table.rows[0].cells[2].text = "AI Agent Action"
    carc_table.rows[0].cells[3].text = "Transmission Channel"

    carc_data = [
        ("CO-16", "Claim lacks required information / medical documentation", "Appeal with itemized statement, certified charts, operative notes", "Payer Portal"),
        ("CO-50", "Non-covered service / Not deemed medically necessary", "Appeal with physician letter & request clinical peer-to-peer review", "Phone / Medical Review"),
        ("CO-197", "Precertification / Prior authorization absent", "Retroactive authorization appeal with emergency/clinical proof", "Secure Portal"),
        ("CO-29", "Timely filing limit exceeded", "Timely filing dispute attaching clearinghouse EDI 999/277 receipts", "Certified Mail / EDI"),
        ("CO-18", "Exact duplicate claim / service disallowance", "Resubmission with procedure modifiers (-59 distinct / -76 repeat)", "EDI Clearinghouse"),
        ("PR-1 / PR-2", "Patient Deductible / Co-insurance responsibility", "Reallocate balance to patient ledger; submit to secondary payer", "Billing System"),
    ]

    for idx, (c, d, a, ch) in enumerate(carc_data, start=1):
        carc_table.rows[idx].cells[0].text = c
        carc_table.rows[idx].cells[1].text = d
        carc_table.rows[idx].cells[2].text = a
        carc_table.rows[idx].cells[3].text = ch

    style_table(carc_table)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 7: TECHNICAL STACK & QUICKSTART COMMANDS
    # -------------------------------------------------------------
    h1 = doc.add_heading("7. Technical Stack & Execution Commands", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    add_callout(
        doc,
        "System Technology Stack Summary",
        "• Frontend: React 19, TypeScript, Vite, Tailwind CSS, Recharts, Lucide Icons\n"
        "• Backend: Python 3.12, FastAPI, Uvicorn, Pydantic v2, SQLAlchemy 2.0 ORM\n"
        "• Data Store: SQLite (backend/data/ar_followup.db), 10,000 Claims\n"
        "• Machine Learning: Scikit-Learn (HistGradientBoostingClassifier), NumPy, Pandas, Joblib\n"
        "• Vector Search (RAG): Scikit-Learn TfidfVectorizer (N-grams 1-2, sublinear TF), Cosine Similarity\n"
        "• Test Suite: Pytest (58 automated test cases covering Phases 1 through 13)"
    )

    doc.add_heading("7.1 Operational Commands", level=2)
    doc.add_paragraph(
        "• Launch Entire Platform: ./run.sh (Starts backend on port 8000 and frontend on port 5173)\n"
        "• Run Pytest Test Suite: PYTHONPATH=. backend/.venv/bin/pytest backend/tests/\n"
        "• Run CLI Demonstration Script: backend/.venv/bin/python scripts/demo_walkthrough.py\n"
        "• Access Frontend Dashboard: http://localhost:5173\n"
        "• Access Swagger API Docs: http://127.0.0.1:8000/docs"
    )

    # Save document
    out_path = "/home/coded-x/Desktop/AR-followup-automation-system/RecoverAI_Comprehensive_System_Architecture_and_Guide.docx"
    doc.save(out_path)
    print(f"Successfully generated Word Document at: {out_path}")

if __name__ == "__main__":
    build_document()
